#!/usr/bin/env python3
"""
gui_app.py — Thai Address Parser GUI
รับ Excel input → แยกที่อยู่ → Export Excel พร้อมคอลัมภ์ใหม่

รองรับ 2 โหมด:
  • GUI  : เปิดด้วยการดับเบิลคลิก (ไม่มี argument)
  • CLI  : ThaiAddressParser.exe input.xlsx --col "ที่อยู่" [--out output.xlsx]
"""

import sys
import os
import threading
import queue
import argparse
import subprocess
from datetime import datetime
from typing import Optional

# ─── Path setup for PyInstaller bundle ───────────────────────────────────────
if getattr(sys, "frozen", False):
    BUNDLE_DIR = sys._MEIPASS          # type: ignore[attr-defined]
    EXE_DIR    = os.path.dirname(sys.executable)
else:
    BUNDLE_DIR = os.path.dirname(os.path.abspath(__file__))
    EXE_DIR    = BUNDLE_DIR

sys.path.insert(0, BUNDLE_DIR)
os.chdir(EXE_DIR)

# Master CSV ที่ bundle ไว้ใน exe (read-only)
MASTER_CSV = os.path.join(BUNDLE_DIR, "data", "thai_administrative_areas.csv")

# Knowledge DB เก็บใน ~/Documents/ThaiAddressParser/data (writable)
APP_DATA_DIR = os.path.join(
    os.path.expanduser("~"), "Documents", "ThaiAddressParser", "data"
)
os.makedirs(APP_DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(APP_DATA_DIR, "address_knowledge.db")

# ─── Imports ─────────────────────────────────────────────────────────────────
from address_parser import KnowledgeBase, AddressParser, ParsedAddress
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────────────────────────────────────
# Excel helpers
# ─────────────────────────────────────────────────────────────────────────────

PARSED_COLS = [
    ("บ้านเลขที่",     "house_number"),
    ("หมู่บ้าน/โครงการ","village"),
    ("หมู่ที่",        "moo"),
    ("ซอย",           "soi"),
    ("ถนน",           "road"),
    ("ตำบล/แขวง",     "sub_district"),
    ("อำเภอ/เขต",     "district"),
    ("จังหวัด",        "province"),
    ("รหัสไปรษณีย์",   "postal_code"),
    ("ประเทศ",         "country"),
    ("ความมั่นใจ (%)", "__confidence"),
]

HEADER_FILL   = PatternFill("solid", fgColor="2C6E8A")
HEADER_FONT   = Font(bold=True, color="FFFFFF", name="Tahoma", size=11)
NEW_COL_FILL  = PatternFill("solid", fgColor="E8F4F8")
BORDER_THIN   = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def style_header(cell, is_new: bool = False):
    cell.fill   = PatternFill("solid", fgColor="1A5276" if is_new else "2C6E8A")
    cell.font   = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER_THIN


def style_data(cell, is_new: bool = False, confidence: float = 1.0):
    if is_new:
        cell.fill = NEW_COL_FILL
    if is_new and "__conf" in str(cell.column_letter):
        pass  # handled separately
    cell.border = BORDER_THIN
    cell.alignment = Alignment(vertical="center")


def get_sheet_names(path: str):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def get_column_names(path: str, sheet: str):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    wb.close()
    return [str(h) for h in headers if h is not None]


def process_excel(
    input_path: str,
    sheet_name: str,
    addr_col: str,
    output_path: str,
    parser: AddressParser,
    progress_cb=None,       # callback(current, total)
    log_cb=None,            # callback(message)
    auto_learn: bool = True,         # บันทึก ≥75% confidence ลง KB อัตโนมัติ
    review_threshold: float = 0.75,  # < threshold → ส่งไป Review sheet
) -> dict:
    """
    อ่าน Excel → parse ทุกแถว → เขียน Excel ใหม่

    การเรียนรู้อัตโนมัติ (auto_learn=True):
      • confidence = 100%  → save as verified=True  (จำแน่นอน ใช้ครั้งต่อไปทันที)
      • confidence 75–99%  → save as verified=False (เก็บ pattern ไว้)
      • confidence < 75%   → ไม่บันทึก + ส่งเข้า "ตรวจสอบ" sheet

    Returns dict with keys: total, ok, learned, uncertain, elapsed, output
    """
    if log_cb: log_cb(f"อ่านไฟล์: {os.path.basename(input_path)}")

    wb_in  = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    ws_in  = wb_in[sheet_name]
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = sheet_name[:31]

    # ── หาคอลัมภ์ที่อยู่ ──
    headers_in = [c.value for c in next(ws_in.iter_rows(max_row=1))]
    col_idx    = None
    for i, h in enumerate(headers_in):
        if str(h).strip() == addr_col.strip():
            col_idx = i
            break
    if col_idx is None:
        wb_in.close()
        raise ValueError(f"ไม่พบคอลัมภ์ '{addr_col}' ใน sheet '{sheet_name}'")

    # ── Main sheet header ──
    out_headers       = [str(h) if h else "" for h in headers_in]
    parsed_hdr_names  = [t for t, _ in PARSED_COLS]
    all_headers       = out_headers + parsed_hdr_names
    ws_out.row_dimensions[1].height = 30
    for ci, h in enumerate(all_headers, 1):
        style_header(ws_out.cell(row=1, column=ci, value=h),
                     is_new=(ci > len(out_headers)))

    # ── Process rows ──
    all_rows  = list(ws_in.iter_rows(min_row=2, values_only=True))
    total     = len(all_rows)
    ok_cnt    = 0
    learned   = 0      # จำนวนที่บันทึกลง KB
    uncertain_rows = []  # (original_row_tuple, ParsedAddress)
    t_start   = datetime.now()

    if log_cb: log_cb(f"พบ {total:,} แถว — เริ่มแยกที่อยู่...")

    for ri, row in enumerate(all_rows, 2):
        raw  = str(row[col_idx]).strip() if row[col_idx] else ""
        addr = parser.parse(raw, remember=False) if raw else ParsedAddress(raw="")

        if addr.province or addr.sub_district:
            ok_cnt += 1

        # ── Auto-learn: บันทึกลง KB ตาม confidence ──
        if auto_learn and raw:
            conf = addr.confidence
            if conf >= review_threshold:
                # 100% → verified (ใช้ exact-match cache ครั้งต่อไป)
                # 75-99% → unverified (เก็บ pattern)
                is_verified = (conf == 1.0)
                parser.kb.save_example(
                    raw, addr.to_dict(),
                    verified=is_verified,
                    source="excel_auto",
                )
                learned += 1
            else:
                # ไม่มั่นใจ → เก็บสำหรับ review
                uncertain_rows.append((row, addr))

        # ── เขียน main sheet ──
        for ci, val in enumerate(row, 1):
            ws_out.cell(row=ri, column=ci, value=val).border = BORDER_THIN

        base = len(out_headers)
        for oi, (_, field) in enumerate(PARSED_COLS):
            ci   = base + oi + 1
            cell = ws_out.cell(row=ri, column=ci)
            if field == "__confidence":
                pct = round(addr.confidence * 100)
                cell.value = pct
                cell.fill  = PatternFill("solid", fgColor=(
                    "D5F5E3" if pct >= 75 else
                    "FCF3CF" if pct >= 40 else "FADBD8"
                ))
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.value = getattr(addr, field, None) or ""
                cell.fill  = NEW_COL_FILL
            cell.border = BORDER_THIN

        if progress_cb and (ri % 50 == 0 or ri - 1 == total):
            progress_cb(ri - 1, total)

    # ── Column widths (main sheet) ──
    for ci in range(1, len(all_headers) + 1):
        ws_out.column_dimensions[get_column_letter(ci)].width = (
            18 if ci > len(out_headers) else 22
        )
    ws_out.freeze_panes = "A2"

    # ── Review sheet (แถวที่ไม่มั่นใจ) ──
    REVIEW_SHEET = "ตรวจสอบ"
    review_editable = [t for t, _ in PARSED_COLS if t != "ความมั่นใจ (%)"]
    review_col_names = [t + " [แก้ไข]" for t in review_editable]

    if uncertain_rows:
        ws_rv = wb_out.create_sheet(REVIEW_SHEET)
        # Header: คอลัมภ์เดิม + คอลัมภ์ที่ parser แยกได้ + คอลัมภ์แก้ไข
        rv_headers = (out_headers
                      + [t for t, _ in PARSED_COLS]
                      + review_col_names)
        ws_rv.row_dimensions[1].height = 30
        for ci, h in enumerate(rv_headers, 1):
            is_edit = ci > len(out_headers) + len(PARSED_COLS)
            cell = ws_rv.cell(row=1, column=ci, value=h)
            cell.fill = PatternFill("solid", fgColor=(
                "C0392B" if is_edit else
                "1A5276" if ci > len(out_headers) else "2C6E8A"
            ))
            cell.font   = HEADER_FONT
            cell.border = BORDER_THIN
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)

        EDIT_FILL = PatternFill("solid", fgColor="FDFEFE")
        EDIT_BORDER = Border(
            left=Side(style="medium", color="C0392B"),
            right=Side(style="medium", color="C0392B"),
            top=Side(style="thin", color="C0392B"),
            bottom=Side(style="thin", color="C0392B"),
        )
        for ri2, (orig_row, addr) in enumerate(uncertain_rows, 2):
            # original data
            for ci, val in enumerate(orig_row, 1):
                ws_rv.cell(row=ri2, column=ci, value=val).border = BORDER_THIN
            # parser result
            base2 = len(out_headers)
            for oi, (_, field) in enumerate(PARSED_COLS):
                ci = base2 + oi + 1
                cell = ws_rv.cell(row=ri2, column=ci)
                if field == "__confidence":
                    cell.value = round(addr.confidence * 100)
                    cell.fill  = PatternFill("solid", fgColor="FADBD8")
                    cell.alignment = Alignment(horizontal="center",
                                               vertical="center")
                else:
                    cell.value = getattr(addr, field, None) or ""
                    cell.fill  = NEW_COL_FILL
                cell.border = BORDER_THIN
            # editable correction columns (ว่าง ให้ user กรอก)
            edit_base = base2 + len(PARSED_COLS)
            for oi in range(len(review_editable)):
                cell = ws_rv.cell(row=ri2, column=edit_base + oi + 1, value="")
                cell.fill   = EDIT_FILL
                cell.border = EDIT_BORDER

        # column widths
        for ci in range(1, len(rv_headers) + 1):
            ws_rv.column_dimensions[get_column_letter(ci)].width = (
                20 if ci > len(out_headers) else 22
            )
        ws_rv.freeze_panes = "A2"

        # คำแนะนำ
        ws_rv.cell(row=1, column=len(rv_headers) + 2,
                   value="💡 กรอกค่าที่ถูกต้องในคอลัมภ์สีแดง แล้วกลับมากด 'เรียนรู้จากการตรวจสอบ'")

    # ── Summary sheet ──
    elapsed = (datetime.now() - t_start).total_seconds()
    pct_ok  = ok_cnt / total * 100 if total else 0
    ws_sum  = wb_out.create_sheet("สรุปผล")
    for row_data in [
        ("ไฟล์ input",          os.path.basename(input_path)),
        ("Sheet",               sheet_name),
        ("คอลัมภ์ที่อยู่",       addr_col),
        ("วันที่ประมวลผล",       datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("จำนวนแถวทั้งหมด",     total),
        ("แยกสำเร็จ",           f"{ok_cnt:,} ({pct_ok:.1f}%)"),
        ("บันทึกลง KB อัตโนมัติ", f"{learned:,} แถว (confidence ≥{review_threshold*100:.0f}%)"),
        ("รอตรวจสอบ",           f"{len(uncertain_rows):,} แถว (ดูใน sheet '{REVIEW_SHEET}')"),
        ("เวลาที่ใช้",           f"{elapsed:.1f}s ({total/elapsed if elapsed else 0:.0f} rows/s)"),
    ]:
        ws_sum.append(row_data)
    ws_sum.column_dimensions["A"].width = 25
    ws_sum.column_dimensions["B"].width = 40

    wb_in.close()
    wb_out.save(output_path)

    return {
        "total":     total,
        "ok":        ok_cnt,
        "learned":   learned,
        "uncertain": len(uncertain_rows),
        "elapsed":   elapsed,
        "output":    output_path,
    }


# ─────────────────────────────────────────────────────────────────────────────
# เรียนรู้จาก Review sheet ที่ user แก้ไขแล้ว
# ─────────────────────────────────────────────────────────────────────────────

REVIEW_SHEET_NAME = "ตรวจสอบ"
REVIEW_FIELD_MAP  = {t + " [แก้ไข]": f for t, f in [
    ("บ้านเลขที่", "house_number"),
    ("หมู่บ้าน/โครงการ", "village"),
    ("หมู่ที่", "moo"),
    ("ซอย", "soi"),
    ("ถนน", "road"),
    ("ตำบล/แขวง", "sub_district"),
    ("อำเภอ/เขต", "district"),
    ("จังหวัด", "province"),
    ("รหัสไปรษณีย์", "postal_code"),
    ("ประเทศ", "country"),
]}


def learn_from_review(
    review_path: str,
    addr_col: str,
    parser: AddressParser,
    log_cb=None,
) -> dict:
    """
    อ่าน Excel ที่มี 'ตรวจสอบ' sheet — บันทึกแถวที่ user กรอกแก้ไขลง KB
    เฉพาะแถวที่มีการกรอกค่าแก้ไขอย่างน้อย 1 คอลัมภ์

    Returns: {"imported": n, "skipped": n}
    """
    wb = openpyxl.load_workbook(review_path, read_only=True, data_only=True)
    if REVIEW_SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(f"ไม่พบ sheet '{REVIEW_SHEET_NAME}' ในไฟล์")

    ws = wb[REVIEW_SHEET_NAME]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) if h else "" for h in rows[0]]

    # หาตำแหน่งคอลัมภ์ที่อยู่ต้นฉบับ และคอลัมภ์แก้ไข
    raw_col_idx  = next((i for i, h in enumerate(headers) if h == addr_col), None)
    edit_col_map = {
        field: i
        for i, h in enumerate(headers)
        if h in REVIEW_FIELD_MAP
        for field in [REVIEW_FIELD_MAP[h]]
    }

    if raw_col_idx is None:
        wb.close()
        raise ValueError(f"ไม่พบคอลัมภ์ '{addr_col}' ใน Review sheet")

    imported = 0
    skipped  = 0

    for row in rows[1:]:
        raw = str(row[raw_col_idx]).strip() if row[raw_col_idx] else ""
        if not raw:
            skipped += 1
            continue

        # เก็บค่าที่ user กรอกแก้ไข
        corrections = {
            field: str(row[ci]).strip()
            for field, ci in edit_col_map.items()
            if ci < len(row) and row[ci] not in (None, "", "None")
        }
        if not corrections:
            skipped += 1
            continue   # user ไม่ได้กรอก → ข้ามไป

        # parse ก่อน แล้ว override ด้วยค่าที่ user แก้ไข
        auto = parser.parse(raw, remember=False)
        corrected_data = auto.to_dict()
        corrected_data.update(corrections)

        corrected = ParsedAddress.from_dict({**corrected_data, "raw": raw})
        corrected.confidence = 1.0

        parser.kb.save_correction(raw, auto.to_dict(), corrected.to_dict())
        if log_cb:
            log_cb(f"✅ เรียนรู้: {raw[:45]}...")
        imported += 1

    wb.close()
    return {"imported": imported, "skipped": skipped}


# ─────────────────────────────────────────────────────────────────────────────
# GUI
# ─────────────────────────────────────────────────────────────────────────────

def build_gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    # ── Init parser (shared) ──
    kb     = KnowledgeBase(DB_PATH)
    parser = AddressParser(kb)

    root = tk.Tk()
    root.title("🏠 Thai Address Parser")
    root.resizable(True, True)
    root.minsize(600, 560)

    try:
        root.tk.call("tk", "scaling", 1.4)
    except Exception:
        pass

    FONT_LABEL  = ("Tahoma", 11)
    FONT_BOLD   = ("Tahoma", 11, "bold")
    FONT_MONO   = ("Courier New", 10)
    BG          = "#F0F4F8"
    ACCENT      = "#2C6E8A"
    BTN_GREEN   = "#27AE60"

    root.configure(bg=BG)
    style = ttk.Style(root)
    style.theme_use("clam")
    style.configure("TFrame",        background=BG)
    style.configure("TLabel",        background=BG, font=FONT_LABEL)
    style.configure("TLabelframe",   background=BG, font=FONT_BOLD)
    style.configure("TLabelframe.Label", background=BG, font=FONT_BOLD, foreground=ACCENT)
    style.configure("TCombobox",     font=FONT_LABEL)
    style.configure("TEntry",        font=FONT_LABEL)
    style.configure("Green.TButton", font=FONT_BOLD, foreground="white", background=BTN_GREEN,
                    padding=(10, 6))
    style.map("Green.TButton",
              background=[("active", "#1E8449"), ("disabled", "#AAB7B8")])
    style.configure("Accent.TButton", font=FONT_LABEL, padding=(6, 4))
    style.configure("TProgressbar",  troughcolor="#D5D8DC", background=ACCENT, thickness=18)

    # ── Variables ──
    var_input   = tk.StringVar()
    var_sheet   = tk.StringVar()
    var_col     = tk.StringVar()
    var_output  = tk.StringVar()
    var_progress= tk.StringVar(value="")
    var_status  = tk.StringVar(value="พร้อมใช้งาน")
    progress_val= tk.IntVar(value=0)
    q           = queue.Queue()

    # ── Layout ──
    main = ttk.Frame(root, padding=16)
    main.pack(fill="both", expand=True)

    # Title
    ttk.Label(main, text="🏠  Thai Address Parser",
              font=("Tahoma", 16, "bold"), foreground=ACCENT).pack(anchor="w", pady=(0, 12))

    # ── Section: Input ──
    frm_in = ttk.LabelFrame(main, text="  📁  ไฟล์ Excel Input  ", padding=(10, 8))
    frm_in.pack(fill="x", pady=(0, 8))

    def browse_input():
        path = filedialog.askopenfilename(
            title="เลือกไฟล์ Excel",
            filetypes=[("Excel files", "*.xlsx *.xlsm *.xls"), ("All files", "*.*")]
        )
        if not path:
            return
        var_input.set(path)
        # สร้าง output path อัตโนมัติ
        base, ext = os.path.splitext(path)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        var_output.set(f"{base}_parsed_{ts}.xlsx")
        # โหลด sheets
        try:
            sheets = get_sheet_names(path)
            cb_sheet["values"] = sheets
            cb_sheet.set(sheets[0])
            on_sheet_changed()
        except Exception as e:
            messagebox.showerror("ข้อผิดพลาด", f"ไม่สามารถอ่านไฟล์ได้:\n{e}")

    def on_sheet_changed(event=None):
        path  = var_input.get()
        sheet = var_sheet.get()
        if not path or not sheet:
            return
        try:
            cols = get_column_names(path, sheet)
            cb_col["values"] = cols
            if cols:
                # ลอง detect คอลัมภ์ที่อยู่อัตโนมัติ
                keywords = ["ที่อยู่", "address", "addr", "full", "ที่อยุ่"]
                found = next(
                    (c for c in cols if any(k in c.lower() for k in keywords)),
                    cols[0]
                )
                cb_col.set(found)
        except Exception:
            pass

    # Row: file path
    row1 = ttk.Frame(frm_in)
    row1.pack(fill="x", pady=2)
    ttk.Entry(row1, textvariable=var_input, width=55).pack(side="left", fill="x", expand=True)
    ttk.Button(row1, text="เลือกไฟล์…", command=browse_input,
               style="Accent.TButton").pack(side="left", padx=(6, 0))

    # Row: sheet
    row2 = ttk.Frame(frm_in)
    row2.pack(fill="x", pady=4)
    ttk.Label(row2, text="Sheet:", width=18, anchor="w").pack(side="left")
    cb_sheet = ttk.Combobox(row2, textvariable=var_sheet, state="readonly", width=30)
    cb_sheet.pack(side="left")
    cb_sheet.bind("<<ComboboxSelected>>", on_sheet_changed)

    # Row: column
    row3 = ttk.Frame(frm_in)
    row3.pack(fill="x", pady=2)
    ttk.Label(row3, text="คอลัมภ์ที่อยู่เต็ม:", width=18, anchor="w").pack(side="left")
    cb_col = ttk.Combobox(row3, textvariable=var_col, width=30)
    cb_col.pack(side="left")
    ttk.Label(row3, text="  (พิมพ์หรือเลือกจากรายการ)", foreground="#888").pack(side="left")

    # ── Section: Output ──
    frm_out = ttk.LabelFrame(main, text="  💾  ไฟล์ Excel Output  ", padding=(10, 8))
    frm_out.pack(fill="x", pady=(0, 8))

    def browse_output():
        path = filedialog.asksaveasfilename(
            title="บันทึกผลลัพธ์",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if path:
            var_output.set(path)

    row4 = ttk.Frame(frm_out)
    row4.pack(fill="x", pady=2)
    ttk.Entry(row4, textvariable=var_output, width=55).pack(side="left", fill="x", expand=True)
    ttk.Button(row4, text="เลือก…", command=browse_output,
               style="Accent.TButton").pack(side="left", padx=(6, 0))

    # ── Run Button ──
    frm_btn = ttk.Frame(main)
    frm_btn.pack(fill="x", pady=(4, 8))

    btn_run = ttk.Button(frm_btn, text="🚀  เริ่มแยกที่อยู่",
                         style="Green.TButton", command=lambda: start_processing())
    btn_run.pack(side="left", ipadx=20)

    btn_open = ttk.Button(frm_btn, text="📂  เปิดไฟล์ผลลัพธ์",
                          style="Accent.TButton", command=lambda: open_output(),
                          state="disabled")
    btn_open.pack(side="left", padx=(10, 0))

    btn_learn = ttk.Button(frm_btn, text="🧠  เรียนรู้จากการตรวจสอบ",
                           style="Accent.TButton", command=lambda: start_learn_from_review(),
                           state="normal")
    btn_learn.pack(side="left", padx=(10, 0))

    # ── Progress ──
    frm_prog = ttk.Frame(main)
    frm_prog.pack(fill="x", pady=(0, 6))
    pb = ttk.Progressbar(frm_prog, variable=progress_val,
                         maximum=100, style="TProgressbar")
    pb.pack(fill="x")
    ttk.Label(frm_prog, textvariable=var_progress,
              foreground="#555", font=("Tahoma", 10)).pack(anchor="e")

    # ── Log ──
    frm_log = ttk.LabelFrame(main, text="  📋  Log  ", padding=(6, 4))
    frm_log.pack(fill="both", expand=True)

    log_text = tk.Text(frm_log, height=10, font=FONT_MONO,
                       state="disabled", wrap="word",
                       bg="#1C2833", fg="#ECF0F1",
                       insertbackground="white", relief="flat")
    log_sb   = ttk.Scrollbar(frm_log, command=log_text.yview)
    log_text.configure(yscrollcommand=log_sb.set)
    log_sb.pack(side="right", fill="y")
    log_text.pack(fill="both", expand=True)

    # Status bar
    status_bar = ttk.Label(main, textvariable=var_status,
                           foreground="#666", font=("Tahoma", 9))
    status_bar.pack(anchor="w", pady=(4, 0))

    # ── Helpers ──
    def log(msg: str):
        ts = datetime.now().strftime("%H:%M:%S")
        log_text.configure(state="normal")
        log_text.insert("end", f"[{ts}]  {msg}\n")
        log_text.see("end")
        log_text.configure(state="disabled")

    def open_output():
        path = var_output.get()
        if path and os.path.exists(path):
            if sys.platform == "darwin":
                subprocess.call(["open", path])
            elif sys.platform == "win32":
                os.startfile(path)
            else:
                subprocess.call(["xdg-open", path])

    # ── Processing ──
    def start_processing():
        input_path  = var_input.get().strip()
        sheet       = var_sheet.get().strip()
        col         = var_col.get().strip()
        output_path = var_output.get().strip()

        if not input_path:
            messagebox.showwarning("⚠️", "กรุณาเลือกไฟล์ Excel ก่อน")
            return
        if not col:
            messagebox.showwarning("⚠️", "กรุณาระบุชื่อคอลัมภ์ที่อยู่")
            return
        if not output_path:
            messagebox.showwarning("⚠️", "กรุณาระบุ path ไฟล์ output")
            return

        btn_run.configure(state="disabled")
        btn_open.configure(state="disabled")
        progress_val.set(0)
        var_progress.set("")
        var_status.set("กำลังประมวลผล...")
        log(f"▶ เริ่ม | Input: {os.path.basename(input_path)}")
        log(f"   Sheet: {sheet} | คอลัมภ์: {col}")

        def _run():
            def _prog(cur, total):
                q.put(("progress", cur, total))
            def _log(msg):
                q.put(("log", msg))
            try:
                result = process_excel(
                    input_path, sheet, col, output_path,
                    parser, _prog, _log,
                    auto_learn=True, review_threshold=0.75,
                )
                q.put(("done", result))
            except Exception as ex:
                q.put(("error", str(ex)))

        threading.Thread(target=_run, daemon=True).start()
        root.after(100, _poll_queue)

    # ── Learn from Review ──
    def start_learn_from_review():
        path = filedialog.askopenfilename(
            title="เลือกไฟล์ Excel ที่ตรวจสอบแล้ว",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")]
        )
        if not path:
            return
        col = var_col.get().strip()
        if not col:
            messagebox.showwarning("⚠️", "กรุณาระบุชื่อคอลัมภ์ที่อยู่ก่อน")
            return

        def _run():
            try:
                result = learn_from_review(path, col, parser, log_cb=lambda m: q.put(("log", m)))
                q.put(("learn_done", result))
            except Exception as ex:
                q.put(("error", str(ex)))

        log(f"🧠 เรียนรู้จาก: {os.path.basename(path)}")
        threading.Thread(target=_run, daemon=True).start()
        root.after(100, _poll_queue)

    def _poll_queue():
        while not q.empty():
            item = q.get_nowait()
            kind = item[0]
            if kind == "progress":
                _, cur, total = item
                pct = int(cur / total * 100) if total else 0
                progress_val.set(pct)
                var_progress.set(f"{cur:,}/{total:,} แถว  ({pct}%)")
            elif kind == "log":
                log(item[1])
            elif kind == "done":
                result = item[1]
                pct_ok = result["ok"] / result["total"] * 100 if result["total"] else 0
                log(f"✅ เสร็จสิ้น {result['total']:,} แถว")
                log(f"   สำเร็จ   : {result['ok']:,} ({pct_ok:.1f}%)"
                    f"  |  เวลา: {result['elapsed']:.1f}s")
                log(f"   บันทึก KB: {result['learned']:,} แถว (auto-learn)")
                if result["uncertain"] > 0:
                    log(f"   ⚠️  รอตรวจสอบ: {result['uncertain']:,} แถว"
                        f" → ดู sheet 'ตรวจสอบ' แล้วกด 'เรียนรู้จากการตรวจสอบ'")
                progress_val.set(100)
                var_progress.set(f"เสร็จแล้ว ✅  {result['total']:,} แถว")
                var_status.set("✅ เสร็จสิ้น")
                btn_run.configure(state="normal")
                btn_open.configure(state="normal")

                review_note = (
                    f"\n⚠️  มี {result['uncertain']:,} แถวที่ไม่มั่นใจ\n"
                    f"   → ดูใน sheet 'ตรวจสอบ' แก้ไข แล้วกด 'เรียนรู้จากการตรวจสอบ'"
                    if result["uncertain"] > 0 else ""
                )
                messagebox.showinfo(
                    "✅ เสร็จสิ้น",
                    f"แยกที่อยู่เสร็จแล้ว!\n\n"
                    f"รายการทั้งหมด : {result['total']:,} แถว\n"
                    f"สำเร็จ        : {result['ok']:,} ({pct_ok:.1f}%)\n"
                    f"บันทึกลง KB   : {result['learned']:,} แถว\n"
                    f"เวลา          : {result['elapsed']:.1f} วินาที"
                    f"{review_note}"
                )
            elif kind == "learn_done":
                result = item[1]
                log(f"🧠 เรียนรู้เสร็จ: {result['imported']} แถว"
                    f" | ข้าม {result['skipped']} แถว (ไม่มีการแก้ไข)")
                stats = kb.get_stats()
                log(f"   KB ตอนนี้: verified={stats['verified_examples']:,}"
                    f" | corrections={stats['corrections']:,}")
                var_status.set(f"🧠 เรียนรู้แล้ว {result['imported']} แถว")
                messagebox.showinfo(
                    "🧠 เรียนรู้เสร็จแล้ว",
                    f"บันทึกลง Knowledge Base แล้ว!\n\n"
                    f"เรียนรู้    : {result['imported']:,} แถว\n"
                    f"ข้าม       : {result['skipped']:,} แถว (ไม่มีการแก้ไข)\n\n"
                    f"KB verified : {stats['verified_examples']:,} รายการ\n"
                    f"รอบต่อไปจะแม่นขึ้นสำหรับที่อยู่รูปแบบนี้ ✅"
                )
            elif kind == "error":
                log(f"❌ Error: {item[1]}")
                var_status.set("❌ เกิดข้อผิดพลาด")
                btn_run.configure(state="normal")
                messagebox.showerror("❌ ข้อผิดพลาด", item[1])

        # Check if still running
        if var_status.get() in ("กำลังประมวลผล...",):
            root.after(150, _poll_queue)

    log("🏠 Thai Address Parser พร้อมใช้งาน")
    log(f"   Knowledge DB: {DB_PATH}")
    stats = kb.get_stats()
    log(f"   Admin areas: {stats.get('administrative_areas', 0):,} ตำบล | "
        f"Keywords: {stats.get('keyword_patterns', 0)}")
    log("─" * 50)

    root.protocol("WM_DELETE_WINDOW", lambda: (kb.close(), root.destroy()))
    root.mainloop()


# ─────────────────────────────────────────────────────────────────────────────
# CLI mode
# ─────────────────────────────────────────────────────────────────────────────

def run_cli():
    ap = argparse.ArgumentParser(
        description="Thai Address Parser — CLI mode",
        formatter_class=argparse.RawTextHelpFormatter,
    )
    ap.add_argument("input",  help="Path ไฟล์ Excel (.xlsx)")
    ap.add_argument("--col",  required=True, metavar="COLUMN",
                    help="ชื่อคอลัมภ์ที่มีที่อยู่เต็ม")
    ap.add_argument("--sheet", default=None,
                    help="ชื่อ Sheet (default: sheet แรก)")
    ap.add_argument("--out",  default=None, metavar="OUTPUT",
                    help="Path ไฟล์ output (default: <input>_parsed_<ts>.xlsx)")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        print(f"❌ ไม่พบไฟล์: {args.input}", file=sys.stderr)
        sys.exit(1)

    # Output path
    if args.out:
        out_path = args.out
    else:
        base = os.path.splitext(args.input)[0]
        ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = f"{base}_parsed_{ts}.xlsx"

    # Sheet name
    sheets = get_sheet_names(args.input)
    sheet  = args.sheet or sheets[0]
    if sheet not in sheets:
        print(f"❌ ไม่พบ sheet '{sheet}' (มี: {', '.join(sheets)})", file=sys.stderr)
        sys.exit(1)

    kb     = KnowledgeBase(DB_PATH)
    parser = AddressParser(kb)

    def progress(cur, total):
        pct = int(cur / total * 100)
        bar = "█" * (pct // 5) + "░" * (20 - pct // 5)
        print(f"\r  [{bar}] {cur:,}/{total:,} ({pct}%)", end="", flush=True)

    def log_fn(msg):
        print(f"  {msg}")

    print(f"\n🏠 Thai Address Parser — CLI Mode")
    print(f"  Input : {args.input}")
    print(f"  Sheet : {sheet}")
    print(f"  Column: {args.col}")
    print(f"  Output: {out_path}\n")

    try:
        result = process_excel(
            args.input, sheet, args.col, out_path,
            parser, progress, log_fn
        )
        print()
        pct_ok = result["ok"] / result["total"] * 100 if result["total"] else 0
        print(f"\n✅ เสร็จสิ้น!")
        print(f"   รายการทั้งหมด  : {result['total']:,}")
        print(f"   สำเร็จ         : {result['ok']:,} ({pct_ok:.1f}%)")
        print(f"   เวลาที่ใช้      : {result['elapsed']:.1f}s")
        print(f"   Output file    : {out_path}")
    except Exception as e:
        print(f"\n❌ Error: {e}", file=sys.stderr)
        sys.exit(1)
    finally:
        kb.close()


# ─────────────────────────────────────────────────────────────────────────────
# Entry Point
# ─────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    if len(sys.argv) > 1:
        run_cli()
    else:
        build_gui()
